#!/usr/bin/env python3
"""
Create a test docx file with sample content
"""

from docx import Document

def create_test_docx():
    # Create a new document
    doc = Document()
    
    # Add a title
    doc.add_heading('Sample Document', 0)
    
    # Add some paragraphs with personal information
    doc.add_paragraph('Hello, my name is John Smith and I work at Microsoft Corporation.')
    doc.add_paragraph('You can contact me at john.smith@microsoft.com or call me at 555-123-4567.')
    doc.add_paragraph('I live in New York City and my birthday is 12/25/1985.')
    doc.add_paragraph('My credit card number is 1234-5678-9012-3456.')
    
    # Add a table with personal data
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'Email'
    table.cell(0, 2).text = 'Phone'
    
    table.cell(1, 0).text = 'John Smith'
    table.cell(1, 1).text = 'john.smith@microsoft.com'
    table.cell(1, 2).text = '555-123-4567'
    
    table.cell(2, 0).text = 'Jane Doe'
    table.cell(2, 1).text = 'jane.doe@google.com'
    table.cell(2, 2).text = '555-987-6543'
    
    # Save the document
    doc.save('media/test_document.docx')
    print("Created test document: media/test_document.docx")

if __name__ == "__main__":
    create_test_docx() 